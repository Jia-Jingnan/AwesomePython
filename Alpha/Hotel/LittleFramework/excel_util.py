# Author:   jingnan
# Date:    2021/6/22
# Desc:

from openpyxl import load_workbook

# 抽象成一个类
class ExcelUtil():

    def __init__(self, file, sheet_name):
        self.file = file
        self.sheet_name = sheet_name

    # 调用该方法获得数据后注意数据类型，需要使用eval()函数转换为原来的数据类型
    def get_param(self):
        workbook = load_workbook(self.file)

        # 定位sheet页
        sheet = workbook[self.sheet_name]

        # 错误演示1:只定位到单元格，并未取值
        # res = sheet.cell(1,1)
        # print(res)


        # 存储所有测试用例数据的列表
        params = []

        for i in range(1, sheet.max_row + 1):

            # 定义存储一个测试用例数据的字典
            param = {}

            # res = sheet.cell(i, 1).value
            # 存入字典
            param['url'] = sheet.cell(i, 1).value

            # 存入字典
            param['method'] = sheet.cell(i, 2).value

            # 存入字典
            param['data'] = sheet.cell(i, 3).value

            # 存入字典
            param['excepted'] = sheet.cell(i, 4).value

            # 将所有的测试用例添加到测试用例列表中
            params.append(param)
        return params



# 测试代码
if __name__ == '__main__':

    # 实例化对象
    bravo = ExcelUtil('bravo.xlsx', 'test')
    # 调用函数
    params = bravo.get_param()
    print(params)